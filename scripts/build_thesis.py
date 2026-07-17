#!/usr/bin/env python3
"""Build the compiled full draft (md/txt/docx) and the template-formatted FINAL_THESIS.docx."""
import re, copy
from pathlib import Path

REPO = Path('/home/user/bachelor')
T = REPO / 'thesis'

CHAPTERS = [T/'1_Introduction.md', T/'2_Literature_Review.md', T/'3_Research_Methodology.md',
            T/'4_Results.md', T/'5_Discussion.md', T/'6_Conclusion.md']
REFS = T/'7_References.md'
APP_B = T/'appendices'/'Appendix_Coding_Records.md'
FRONT = T/'0_Front_Matter.md'

def front_parts():
    txt = FRONT.read_text()
    abstract = txt.split('## Abstract')[1].split('## Acknowledgements')[0].strip()
    ack = txt.split('## Acknowledgements')[1].strip()
    return abstract, ack

# ---------- 1) compiled markdown ----------
abstract, ack = front_parts()
md_parts = ['# Barriers to Reshoring from Asian Markets: A Systematic Literature Review\n',
            '## Abstract\n\n' + abstract + '\n',
            '## Acknowledgements\n\n' + ack + '\n']
for ch in CHAPTERS:
    md_parts.append(ch.read_text().strip() + '\n')
md_parts.append(REFS.read_text().strip() + '\n')
md_parts.append(APP_B.read_text().strip() + '\n')
full_md = '\n\n'.join(md_parts)
(T/'COMPILED_DRAFT_FULL.md').write_text(full_md)

# ---------- 2) plain txt (markdown stripped) ----------
def strip_md(s):
    out = []
    for line in s.splitlines():
        l = line
        l = re.sub(r'^#{1,6}\s*', '', l)
        l = l.replace('**', '')
        l = re.sub(r'(?<![\w*])\*([^*\n]+)\*(?![\w*])', r'\1', l)
        l = re.sub(r'^\|[-| ]+\|$', '', l)  # table separator rows
        l = re.sub(r'^\|\s*', '', l); l = re.sub(r'\s*\|$', '', l)
        l = l.replace(' | ', '   ')
        l = l.replace('```', '')
        l = re.sub(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', r'[\1 — image: \2]', l)
        out.append(l)
    return '\n'.join(out)
(T/'COMPILED_DRAFT_FULL.txt').write_text(strip_md(full_md))

# ---------- 3) docx builders ----------
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def base_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for i in (1,2,3):
        h = doc.styles[f'Heading {i}']
        h.font.name = 'Times New Roman'; h.font.color.rgb = None
        h.font.size = Pt(16 if i==1 else 14 if i==2 else 12)
        h.font.bold = True
        h.element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
        h.element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        try:
            h.font.color.rgb = None
        except Exception:
            pass
        # force black
        rPr = h.element.get_or_add_rPr()
        c = rPr.find(qn('w:color'))
        if c is None:
            c = OxmlElement('w:color'); rPr.append(c)
        c.set(qn('w:val'), '000000')
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    return doc

BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
ITAL_RE = re.compile(r'(?<![\w*])\*([^*\n]+?)\*(?![\w*])')

def add_runs(par, text):
    # tokenize bold then italics
    pos = 0
    tokens = []
    pattern = re.compile(r'\*\*(.+?)\*\*|(?<![\w*])\*([^*\n]+?)\*(?![\w*])')
    for m in pattern.finditer(text):
        if m.start() > pos:
            tokens.append(('plain', text[pos:m.start()]))
        if m.group(1) is not None:
            tokens.append(('bold', m.group(1)))
        else:
            tokens.append(('ital', m.group(2)))
        pos = m.end()
    if pos < len(text):
        tokens.append(('plain', text[pos:]))
    for kind, t in tokens:
        r = par.add_run(t)
        if kind == 'bold': r.bold = True
        if kind == 'ital': r.italic = True

def add_table(doc, header, rows, font_pt=10):
    tbl = doc.add_table(rows=1+len(rows), cols=len(header))
    tbl.style = 'Table Grid'
    for j,h in enumerate(header):
        cell = tbl.rows[0].cells[j]
        p = cell.paragraphs[0]; add_runs(p, h)
        for r in p.runs: r.bold = True; r.font.size = Pt(font_pt)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            cell = tbl.rows[i+1].cells[j]
            p = cell.paragraphs[0]; add_runs(p, val)
            for r in p.runs: r.font.size = Pt(font_pt)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return tbl

def parse_md_blocks(text):
    """Yield ('h1'|'h2'|'h3'|'p'|'table'|'code', payload)."""
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1; continue
        if line.startswith('```'):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i]); i += 1
            i += 1
            yield ('code', block); continue
        m_img = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', line.strip())
        if m_img:
            yield ('img', m_img.group(2))
            i += 1; continue
        if line.startswith('|'):
            tbl = []
            while i < len(lines) and lines[i].startswith('|'):
                tbl.append(lines[i]); i += 1
            rows = []
            for tr in tbl:
                if re.match(r'^\|[-| ]+\|$', tr): continue
                cells = [c.strip() for c in tr.strip().strip('|').split('|')]
                rows.append(cells)
            if rows:
                yield ('table', (rows[0], rows[1:]))
            continue
        m = re.match(r'^(#{1,3})\s+(.*)$', line)
        if m:
            yield ('h%d' % len(m.group(1)), m.group(2).strip())
            i += 1; continue
        # paragraph: accumulate until blank
        par = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(('|','#','```','![')):
            par.append(lines[i]); i += 1
        yield ('p', ' '.join(par))

def landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)

def portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    if section.page_width > section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)

CAPTION_RE = re.compile(r'^\*\*Table \d+\.\*\*')

def render_blocks(doc, text, landscape_tables=False, table_font=10):
    pending_caption = None
    for kind, payload in parse_md_blocks(text):
        if kind == 'h1':
            doc.add_heading(payload, level=1)
        elif kind == 'h2':
            doc.add_heading(payload, level=2)
        elif kind == 'h3':
            doc.add_heading(payload, level=3)
        elif kind == 'img':
            img_path = T / payload
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            r = p.add_run()
            r.add_picture(str(img_path), width=Inches(6.2))
        elif kind == 'code':
            for cl in payload:
                p = doc.add_paragraph()
                r = p.add_run(cl); r.font.name = 'Courier New'; r.font.size = Pt(9)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.space_after = Pt(0)
        elif kind == 'table':
            header, rows = payload
            if landscape_tables:
                s = doc.add_section(WD_SECTION.NEW_PAGE); landscape(s)
                if pending_caption:
                    p = doc.add_paragraph(); add_runs(p, pending_caption)
                    pending_caption = None
                add_table(doc, header, rows, font_pt=table_font)
                s2 = doc.add_section(WD_SECTION.NEW_PAGE); portrait(s2)
            else:
                if pending_caption:
                    p = doc.add_paragraph(); add_runs(p, pending_caption)
                    pending_caption = None
                add_table(doc, header, rows, font_pt=table_font)
        else:
            if landscape_tables and CAPTION_RE.match(payload):
                # hold the caption so it lands inside the landscape section with its table
                pending_caption = payload
                continue
            if pending_caption:
                p = doc.add_paragraph(); add_runs(p, pending_caption)
                pending_caption = None
            p = doc.add_paragraph()
            add_runs(p, payload)
    if pending_caption:
        p = doc.add_paragraph(); add_runs(p, pending_caption)

def add_page_numbers(doc):
    # bottom-centre page number on every section; hide on the literal first page (title page)
    sec0 = doc.sections[0]
    sec0.different_first_page_header_footer = True
    for sec in doc.sections:
        footer = sec.footer
        footer.is_linked_to_previous = False if sec is sec0 else True
    fp = sec0.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r'); tt = OxmlElement('w:t'); tt.text = '1'
    r.append(tt); fld.append(r)
    fp._p.append(fld)
    for run in fp.runs:
        run.font.name = 'Times New Roman'; run.font.size = Pt(12)

def add_toc_field(doc):
    p = doc.add_paragraph()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    r = OxmlElement('w:r'); t = OxmlElement('w:t')
    t.text = 'Right-click and choose "Update Field" to generate the Table of Contents.'
    r.append(t); fld.append(r)
    p._p.append(fld)

# ===== FINAL_THESIS.docx =====
doc = base_doc()
# Title page
def cpar(text, size=12, bold=False, before=0, after=6):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    return p

cpar('BACHELOR THESIS', 16, True, before=60, after=30)
cpar('Title of the Bachelor‘s Thesis', 12, False, after=10)
cpar('BARRIERS TO RESHORING FROM ASIAN MARKETS:', 14, True, after=0)
cpar('A SYSTEMATIC LITERATURE REVIEW', 14, True, after=40)
cpar('verfasst von / submitted by', 12, False, after=6)
cpar('Serdiuk Aleksandra', 12, True, after=0)
cpar('12045860', 12, False, after=40)
cpar('Vienna, July 2026', 12, False, after=40)
cpar('Studienrichtung lt. Studienblatt / degree programme as it appears on the student record sheet:', 11, False, after=4)
cpar('Bachelorstudium Internationale Betriebswirtschaft', 12, False, after=20)
cpar('Betreut von / Supervisor: Dr. Aveed Raha', 12, False)
doc.add_page_break()

# Acknowledgments
doc.add_heading('Acknowledgments', level=1)
p = doc.add_paragraph(); add_runs(p, ack)
doc.add_page_break()
# Abstract
doc.add_heading('Abstract', level=1)
p = doc.add_paragraph(); add_runs(p, abstract)
doc.add_page_break()
# ToC
doc.add_heading('Table of Contents', level=1)
add_toc_field(doc)
doc.add_page_break()

for ch in CHAPTERS:
    txt = ch.read_text()
    render_blocks(doc, txt, landscape_tables=False, table_font=10)
    doc.add_page_break()

render_blocks(doc, REFS.read_text())
# Per the supervisor's instruction (17 July 2026): the PRISMA flow diagram (Figure 1)
# and the coding table (Table 1) sit in Chapter 3; the single Appendix carries the
# detailed per-source coding records. The old Appendix A was superseded by Figure 1.
doc.add_page_break()
render_blocks(doc, APP_B.read_text(), table_font=9)

add_page_numbers(doc)
doc.save(str(T/'FINAL_THESIS.docx'))

# ===== COMPILED_DRAFT_FULL.docx (simple clean formatting, no title page/toc) =====
doc2 = base_doc()
render_blocks(doc2, full_md, landscape_tables=False, table_font=10)
doc2.save(str(T/'COMPILED_DRAFT_FULL.docx'))
print('built: COMPILED_DRAFT_FULL.md/.txt/.docx and FINAL_THESIS.docx')
